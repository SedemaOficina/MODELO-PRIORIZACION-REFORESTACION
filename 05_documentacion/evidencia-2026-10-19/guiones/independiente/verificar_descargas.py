# Verifica los archivos descargados (descargas_app.js) contra el recalculo independiente y contra la lista de la pantalla.
# Uso: python -I verificar_descargas.py <carpeta de guiones> <carpeta de descargas>
import sys, json, re, os, subprocess; sys.path.insert(0, sys.argv[1])
import datos_indep as D
import openpyxl
d = D.carga('X:/r7/docs'); F, V, M, EF, EV = d['F'], d['V'], d['M'], d['EF'], d['EV']
BIT = {'cv': 1, 'to': 2, 'am': 4}
def tipos(c): return (1 if c[0] else 0) | (2 if c[1] else 0) | (4 if c[2] else 0)
def parse(q):
    p = dict(x.split('=') for x in q.split('&') if x)
    m = sum(BIT[x] for x in p.get('e', '').split(',') if x)
    return p.get('r', 'alc'), p.get('a'), int(p['c']) if 'c' in p else None, p.get('b', 'todas'), m
def esperado(q):
    r, a, c, b, m = parse(q)
    mun = M['muns'].index(a) if a else (M['muns'].index(M['colonias'][c]['m']) if c else None)
    terr = (lambda f: f['col'] == c) if c else (lambda f: mun is None or f['mun'] == mun)
    pasaB = lambda f: b == 'todas' or ((f['banq'] == 0) == (b == 'con'))
    pasaE = lambda i: m == 0 or (i in EF and tipos(EF[i]) & m)
    pasaEV = lambda j: m == 0 or (j >= 0 and j in EV and tipos(EV[j]) & m)
    fr = [f for f in F if not f['gc'] and terr(f) and pasaB(f) and pasaE(f['i']) and f['prio'] >= 2]
    vp = [v for v in V if (mun is None or v['mun'] == mun) and pasaEV(v['i']) and v['prio'] >= 2]
    gcdoc = [f for f in F if f['gc'] and terr(f) and pasaEV(f['vp']) and f['prio'] >= 2]
    gctab = [f for f in F if f['gc'] and terr(f) and pasaB(f) and pasaE(f['i']) and f['prio'] >= 2]
    return dict(fr=fr, vp=vp, gcdoc=sum(f['len'] for f in gcdoc)/1000, gctab=sum(f['len'] for f in gctab)/1000, colonia=c)
def libro(p):
    wb = openpyxl.load_workbook(p, read_only=True)
    datos = [list(r) for r in wb['Datos'].iter_rows(values_only=True)]
    dic = {r[0]: r[1] for r in wb['Diccionario'].iter_rows(values_only=True) if r and r[0]}
    return datos, dic
out = []
idx = json.load(open(os.path.join(sys.argv[2], 'indice.json'), encoding='utf8'))
for s in idx['res']:
    E = esperado(s['q']); out.append(f"\n== {s['q']} (orden {s['pantalla']['orden']})")
    for a in s['archivos']:
        if 'archivo' not in a: out.append(f"  {a['boton']}: {a['estado']}"); continue
        p = os.path.join(sys.argv[2], a['archivo'])
        if a['archivo'].endswith('.xlsx'):
            datos, dic = libro(p); h = datos[0]; filas = datos[1:]
            if 'frentes_prioritarios' in a['archivo']:
                ids = {r[h.index('id_frente')] for r in filas}; ex = {f['i'] for f in E['fr']}
                out.append(f"  frentes: {len(filas)} renglones · recalculo {len(ex)} · faltan {len(ex-ids)} · sobran {len(ids-ex)}")
                k = [x for x in dic if x and 'Gobierno Central' in str(x) and 'Prioritario' in str(x)]
                for x in k: out.append(f"  diccionario «{x}»: {dic[x]} · regla documentada (sin banqueta, espacio por tramo) {round(E['gcdoc'],2)} · con banqueta y espacio por frente {round(E['gctab'],2)}")
                if 'cepas_vacantes' in h:
                    bad = sum(1 for r in filas if tuple(r[h.index(c)] for c in ('cepas_vacantes','tocones','arboles_muertos')) != EF.get(r[h.index('id_frente')], (0,0,0)))
                    out.append(f"  columnas de espacios distintas del archivo de datos: {bad}")
            elif 'resumen_calles' in a['archivo']:
                nombres = [f"{r[0]}|{r[1]}" for r in filas]
                lista = s['pantalla']['lista']
                out.append(f"  calles: {len(filas)} renglones; pantalla: «{s['pantalla']['iniCount']}», {len(lista)} renglones visibles")
                # orden: el renglón k de la pantalla contra el renglón k del archivo
                dif = []
                for k, t in enumerate(lista[:len(filas)]):
                    nom = filas[k][0]
                    if not re.sub(r'^\d+', '', t).startswith(str(nom)): dif.append((k+1, t[:60], nom, filas[k][1]))
                out.append(f"  renglones de la pantalla que no coinciden con el mismo renglón del Excel: {len(dif)} de {min(len(lista), len(filas))}")
                for x in dif[:6]: out.append(f"    #{x[0]} pantalla «{x[1]}» · Excel «{x[2]} ({x[3]})»")
                if 'cepas_vacantes' in h:
                    out.append(f"  orden del Excel por: {'km_prioritario' if all(filas[i][13] >= filas[i+1][13] for i in range(len(filas)-1)) else 'otro'}")
            elif 'tramos_prioritarios' in a['archivo']:
                ex = len(E['vp']); out.append(f"  tramos: {len(filas)} renglones · recalculo {ex}")
                for x in dic:
                    if x and 'Equivalente' in str(x): out.append(f"  diccionario «{x}»: {dic[x]}")
            elif 'resumen_avenidas' in a['archivo']:
                lista = s['pantalla']['lista']; dif = 0
                for k, t in enumerate(lista[:len(filas)]):
                    if not re.sub(r'^\d+', '', t).startswith(str(filas[k][0])): dif += 1
                out.append(f"  avenidas: {len(filas)} renglones; pantalla «{s['pantalla']['iniCount']}»; renglones en distinto orden: {dif} de {min(len(lista), len(filas))}")
                if dif: out.append(f"    pantalla 1-3: {[re.sub(r'^\\d+', '', t)[:40] for t in lista[:3]]} · Excel 1-3: {[r[0] for r in filas[:3]]}")
        elif a['archivo'].endswith('.geojson'):
            g = json.load(open(p, encoding='utf8')); fs = g['features']
            fr = {x['properties']['id_frente'] for x in fs if 'id_frente' in x['properties']}; nv = sum(1 for x in fs if 'id_tramo' in x['properties'])
            out.append(f"  geojson: {len(fr)} frentes (recalculo {len(E['fr'])}, iguales {fr == {f['i'] for f in E['fr']}}) · {nv} tramos (recalculo {len(E['vp']) if not E['colonia'] else 0})")
        elif a['archivo'].endswith('.kml'):
            t = open(p, encoding='utf8').read(); out.append(f"  kml: {t.count('<Placemark>')} marcas · recalculo {len(E['fr'])}+{len(E['vp'])}")
        elif a['archivo'].endswith('.pdf'):
            txt = subprocess.run(['pdftotext', '-layout', p, '-'], capture_output=True).stdout.decode('utf8', 'replace')
            open(p + '.txt', 'w', encoding='utf8').write(txt)
            g = re.search(r'Gobierno Central: ([\d.,]+ (?:km|m)) de frente prioritario', txt.replace('\n', ' '))
            out.append(f"  pdf: {len(txt)} caracteres" + (f" · «Gobierno Central: {g.group(1)} de frente prioritario» · regla documentada {round(E['gcdoc'],2)} km · con banqueta y espacio por frente {round(E['gctab'],2)} km" if g else ''))
open(os.path.join(sys.argv[2], 'verificacion.txt'), 'w', encoding='utf8').write('\n'.join(out))
