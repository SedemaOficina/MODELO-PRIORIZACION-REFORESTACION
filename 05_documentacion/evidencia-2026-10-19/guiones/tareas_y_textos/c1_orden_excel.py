# C1 · Compara los 10 primeros de pantalla («Calles · más cepas vacantes») con el Excel «Resumen por calle»: como viene y ordenado por
# cepas_vacantes, la única columna de cepas que trae. Correr desde la carpeta de evidencia: python -I - < c1_orden_excel.py
import os, json
from openpyxl import load_workbook
lp = lambda p: ('\\\\?\\' + os.path.abspath(p)) if os.name == 'nt' else p
aqui = os.getcwd()
for caso in json.load(open(lp(os.path.join(aqui, 'c1_orden_excel.json')), encoding='utf-8')):
    wb = load_workbook(lp(os.path.join(aqui, 'descargas', caso['excel'])), read_only=True); f = list(wb['Datos'].iter_rows(values_only=True)); cab = list(f[0]); d = f[1:]
    ic, ik, icol = cab.index('cepas_vacantes'), cab.index('km_prioritario'), cab.index('colonia')
    clave = lambda r: (str(r[0]), str(r[icol]))
    pant = [(x['calle'], x['colonia']) for x in caso['lista']]
    como_viene = [clave(r) for r in d[:10]]
    por_cepas = [clave(r) for r in sorted(d, key=lambda r: (-(r[ic] or 0), -(r[ik] or 0)))[:10]]
    print(f"\n== {caso['alcaldia']} · {caso['conteo']} · {len(d)} renglones en el Excel")
    print('  Pantalla:', [f'{a} ({b})' for a, b in pant])
    print('  Pantalla, cifras:', [x['k'] for x in caso['lista']][:4], '…')
    print('  Excel como viene, coincide en posición:', sum(a == b for a, b in zip(pant, como_viene)), 'de 10; en el conjunto:', len(set(pant) & set(como_viene)))
    print('  Excel ordenado por cepas_vacantes, coincide en posición:', sum(a == b for a, b in zip(pant, por_cepas)), 'de 10; en el conjunto:', len(set(pant) & set(por_cepas)))
    print('  Excel por cepas_vacantes:', [f'{a} ({b})' for a, b in por_cepas])
