# C1 · Compara el orden de «Calles · más cepas vacantes» en pantalla (c1_tareas_<vp>.json) con el Excel «Resumen por calle» descargado.
# Si la ruta del guion pasa de 260 caracteres (Windows), correrlo desde su carpeta con: python -I - tel < c1_excel_contra_pantalla.py
# Uso: python -I c1_excel_contra_pantalla.py tel     (lee descargas/resumen_calles_*.xlsx y la lista de pantalla)
import os, sys, json, glob
from openpyxl import load_workbook
lp = lambda p: ('\\\\?\\' + os.path.abspath(p)) if os.name == 'nt' else p
aqui = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else os.getcwd(); vp = sys.argv[1] if len(sys.argv) > 1 else 'tel'
o = json.load(open(lp(os.path.join(aqui, f'c1_tareas_{vp}.json')), encoding='utf-8'))
A = o['tareas']['A']; arch = os.path.join(aqui, 'descargas', A['archivoCalles'])
wb = load_workbook(lp(arch), read_only=True); ws = wb['Datos']; filas = list(ws.iter_rows(values_only=True)); cab = list(filas[0]); datos = filas[1:]
print('Columnas:', cab)
ic = cab.index('cepas_vacantes'); ik = cab.index('km_prioritario')
print('Renglones:', len(datos), '· orden del archivo (primeros 10, como viene):')
for r in datos[:10]: print('   ', r[0], '| km_prioritario', r[ik], '| cepas_vacantes', r[ic])
pant = [x['calle'] for x in A['listaPantalla']]
xl_por_cepas = sorted(datos, key=lambda r: (-(r[ic] or 0), -(r[ik] or 0)))
top_xl = [str(r[0]) for r in xl_por_cepas[:10]]
print('\nPantalla (Calles · más cepas vacantes):', pant)
print('Excel ordenado por cepas_vacantes (desc) y km_prioritario:', top_xl)
print('Excel tal como viene (primeros 10):', [str(r[0]) for r in datos[:10]])
coinc = sum(1 for a, b in zip(pant, top_xl) if a == b)
print(f'\nCoinciden en la misma posición: {coinc} de 10; en el conjunto: {len(set(pant) & set(top_xl))} de 10')
# el diccionario: qué dice de cepas_vacantes
wd = wb['Diccionario']
for r in wd.iter_rows(values_only=True):
    if r and r[0] and str(r[0]).startswith(('cepas_vacantes', 'Ámbito', 'Filtro', 'Contenido')): print('Diccionario ·', r[0], '·', r[1])
