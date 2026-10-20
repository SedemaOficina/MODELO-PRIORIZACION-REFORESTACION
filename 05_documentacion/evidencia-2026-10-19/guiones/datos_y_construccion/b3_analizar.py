# B2/B3 · Abre cada archivo descargado y lo compara con el recálculo independiente; revisa inyección de fórmulas,
# acentos, metadatos y nombres. Uso: python b3_analizar.py <docs> <carpeta de descargas (con descargas.json)>
import sys, os, json, re, zipfile, subprocess, xml.etree.ElementTree as ET
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from b2_decodificador import carga
from urllib.parse import parse_qs
import openpyxl
D = carga(sys.argv[1]); meta, F, V, P, EF, EV = D; SAL = sys.argv[2]
TIPO = {'cv': 0, 'to': 1, 'am': 2}
res = json.load(open(os.path.join(SAL, 'descargas.json'), encoding='utf-8'))
PELIGRO = re.compile(r'^[=+\-@\t\r]')
def esperado(q):
    r = q.get('r', 'alc'); a = q.get('a'); c = int(q['c']) if 'c' in q else None; b = q.get('b', 'todas'); t = [TIPO[x] for x in q['e'].split(',')] if 'e' in q else []
    mun = meta['muns'].index(a) if a else None
    fr = [i for i, f in enumerate(F) if not f['gc'] and f['prio'] >= 2 and ((f['col'] == c) if c is not None else (mun is None or f['mun'] == mun))
          and (b == 'todas' or (b == 'con') == (f['banq'] == 0)) and (not t or any(EF[1].get(i, (0, 0, 0))[x] for x in t))]
    vmun = mun if c is None else meta['muns'].index(meta['colonias'][c]['m'])
    vp = [k for k, v in enumerate(V) if v['prio'] >= 2 and (vmun is None or v['mun'] == vmun) and (not t or any(EV[1].get(k, (0, 0, 0))[x] for x in t))]
    return r, fr, vp
out = []
for d in res:
    if not d.get('nombre'): continue
    q = {k: v[0] for k, v in parse_qs(d['consulta']).items()}; r, fr, vp = esperado(q); ruta = d['ruta']; n = d['nombre']; nota = []
    if n.endswith('.xlsx'):
        wb = openpyxl.load_workbook(ruta, read_only=True); ws = wb['Datos']; filas = list(ws.iter_rows(values_only=True)); cab = filas[0]; datos = filas[1:]
        malas = [(i, v) for row in datos for i, v in enumerate(row) if isinstance(v, str) and PELIGRO.match(v)]
        dic = [x for x in wb['Diccionario'].iter_rows(values_only=True)]
        mal_dic = [v for row in dic for v in row if isinstance(v, str) and PELIGRO.match(v)]
        props = wb.properties
        nota.append('filas %d' % len(datos))
        if 'id_frente' in cab:
            ids = [row[cab.index('id_frente')] for row in datos]; exp = set(fr) if d['id'] == 'dl-frentes' else None
            if exp is not None: nota.append('id_frente = recálculo: %s (faltan %d, sobran %d)' % (set(ids) == exp, len(exp - set(ids)), len(set(ids) - exp)))
            km = sum(row[cab.index('longitud_m')] for row in datos) / 1000; nota.append('km %.3f vs %.3f' % (km, sum(F[i]['len'] for i in fr) / 1000))
            for k, col in enumerate(('cepas_vacantes', 'tocones', 'arboles_muertos')):
                if col in cab: s = sum(row[cab.index(col)] or 0 for row in datos); e = sum(EF[1].get(i, (0, 0, 0))[k] for i in fr); nota.append('%s %d vs %d' % (col, s, e))
        if 'id_tramo' in cab:
            ids = set(row[cab.index('id_tramo')] for row in datos); exp = set(V[k]['rec'] for k in vp)
            lon = sorted(row[cab.index('longitud_m')] for row in datos) == sorted(V[k]['len'] for k in vp)
            nota.append('id_tramo (registro de origen) = recálculo: %s · renglones %d vs partes %d · longitudes iguales: %s' % (ids == exp, len(datos), len(vp), lon))
            for k, col in enumerate(('cepas_vacantes', 'tocones', 'arboles_muertos')):
                if col in cab: s = sum(row[cab.index(col)] or 0 for row in datos); e = sum(EV[1].get(i, (0, 0, 0))[k] for i in vp); nota.append('%s %d vs %d' % (col, s, e))
        if 'km_prioritario' in cab:
            s = sum(row[cab.index('km_prioritario')] or 0 for row in datos)
            e = (sum(F[i]['len'] for i in fr) if 'frentes_total' in cab else sum(V[k]['len'] for k in vp)) / 1000
            nota.append('suma km_prioritario %.2f vs %.2f' % (s, e))
        nota.append('celdas que empiezan con = + - @: %d en datos, %d en diccionario' % (len(malas), len(mal_dic)))
        nota.append('título «%s» · idioma %s · autor %s' % (props.title, props.language, props.creator))
        z = zipfile.ZipFile(ruta); hoja = z.read('xl/worksheets/sheet1.xml')[:4000].decode('utf-8', 'replace'); nota.append('encabezado inmovilizado: %s' % ('<pane' in hoja))
    elif n.endswith('.geojson'):
        g = json.load(open(ruta, encoding='utf-8')); fs_ = g['features']
        props = [f['properties'] for f in fs_]
        nota.append('features %d · esperado (frentes %d + tramos %d si aplica)' % (len(fs_), len(fr), len(vp)))
        keys = set(k for p in props for k in p); nota.append('campos ' + ','.join(sorted(keys))[:200])
        fuera = sum(1 for f in fs_ for x, y in (f['geometry']['coordinates'] if f['geometry']['type'] == 'LineString' else [c for l in f['geometry']['coordinates'] for c in l]) if not (-99.4 < x < -98.9 and 19.0 < y < 19.6))
        nota.append('vértices fuera de la CDMX %d' % fuera)
        nota.append('metadatos ' + json.dumps({k: v for k, v in g.items() if k != 'features'}, ensure_ascii=False)[:300])
    elif n.endswith('.kml'):
        t = open(ruta, encoding='utf-8').read(); pm = t.count('<Placemark'); nota.append('Placemark %d' % pm)
        try: ET.fromstring(t.encode('utf-8')); nota.append('XML válido')
        except Exception as e: nota.append('XML inválido: %s' % e)
        nota.append('acentos íntegros: %s' % ('Ã' not in t))
    elif n.endswith('.pdf'):
        txt = subprocess.run(['pdftotext', '-enc', 'UTF-8', '-layout', ruta, '-'], capture_output=True).stdout.decode('utf-8', 'replace')
        import pypdf; rd = pypdf.PdfReader(ruta); mi = rd.metadata or {}
        lang = rd.trailer['/Root'].get('/Lang'); mk = rd.trailer['/Root'].get('/MarkInfo')
        nota.append('páginas %d · título «%s» · autor «%s» · idioma %s · etiquetado %s' % (len(rd.pages), mi.get('/Title'), mi.get('/Author'), lang, mk))
        nota.append('mojibake: %s · NaN/undefined: %s' % ('Ã' in txt or '�' in txt, bool(re.search(r'\bNaN\b|undefined|Infinity', txt))))
        open(ruta + '.txt', 'w', encoding='utf-8').write(txt)
    out.append('%s · %s · %s\n    %s' % (d['consulta'], d['id'], n, '\n    '.join(nota)))
print('\n'.join(out))
