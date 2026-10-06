# -*- coding: utf-8 -*-
"""Listas de los catálogos de colonias y de calles para su homologación en el SIA (auditoría H-029 y H-101).

Solo lee los datos publicados (02_fuente/datos/*.bin) y escribe un Excel con las listas. NO modifica ningún catálogo:
la corrección corresponde a la fuente de cada capa.

Uso:  python3 03_procesamiento_datos/reporte_catalogos.py
Sale: 06_entregables/Catalogos_para_homologacion_SIA.xlsx
Requiere: Python 3 con openpyxl.
"""
import gzip
import json
import os
import re
import unicodedata
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATOS = os.path.join(RAIZ, '02_fuente', 'datos')
SALIDA = os.path.join(RAIZ, '06_entregables', 'Catalogos_para_homologacion_SIA.xlsx')


def lector(raw):
    """Enteros varint con signo en zigzag (ver ARQUITECTURA.md, sección 7)."""
    pos = 0
    n = len(raw)

    def rv():
        nonlocal pos
        res = 0
        shift = 0
        while True:
            b = raw[pos]
            pos += 1
            res |= (b & 0x7f) << shift
            shift += 7
            if not b & 0x80:
                break
        return -((res + 1) >> 1) if res & 1 else res >> 1
    rv.fin = lambda: pos >= n
    return rv


META = json.loads(gzip.open(os.path.join(DATOS, 'meta.bin')).read().decode('utf-8'))
rv = lector(gzip.open(os.path.join(DATOS, 'data.bin')).read())
N = rv()
FR = []   # (alcaldía, prioridad, nombre, colonia, longitud en m, a cargo de Gobierno Central, lon, lat del primer vértice)
px = py = 0
for i in range(N):
    mun, prio, name, tipo, col, ln, flags, vp = (rv() for _ in range(8))
    nv = rv()
    x0 = y0 = None
    for k in range(nv):
        px += rv()
        py += rv()
        if k == 0:
            x0, y0 = px, py
    FR.append((mun, prio, name, col, ln, (flags >> 6) & 1, x0 / META['Q'], y0 / META['Q']))
rv = lector(gzip.open(os.path.join(DATOS, 'vp.bin')).read())
NV = rv()
VP = []   # (nomenclatura, nombre, alcaldía, longitud en m)
for i in range(NV):
    nom, nombre, tipo, car, circ, alct, mun, prio, ln, clave, rec = (rv() for _ in range(11))
    nv = rv()
    for k in range(nv):
        rv()
        rv()
    VP.append((nom, nombre, mun, ln))

ALC = dict(zip(META['muns'], META['munNames']))
ALC_I = META['munNames']
COL = META['colonias']
NOM = META['names']


def norm(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower()
    return re.sub(r'[^a-z0-9ñ]+', ' ', s).strip()


# frentes y km por colonia y por nombre de calle
col_n = defaultdict(int)
col_km = defaultdict(float)
nom_n = defaultdict(int)
nom_km = defaultdict(float)
nom_alc = defaultdict(set)
for mun, prio, name, col, ln, gc, x, y in FR:
    col_n[col] += 1
    col_km[col] += ln / 1000
    nom_n[name] += 1
    nom_km[name] += ln / 1000
    nom_alc[name].add(ALC_I[mun])

hojas = []   # (título, descripción, encabezados, renglones)

# ---------- colonias ----------
grupos = defaultdict(list)
for i, c in enumerate(COL):
    if c.get('n'):
        grupos[(c['m'], norm(c['n']))].append(i)
r = []
for (m, _), ids in sorted(grupos.items(), key=lambda kv: (ALC.get(kv[0][0], ''), kv[0][1])):
    if len(ids) < 2:
        continue
    cps = [COL[i]['cp'] for i in ids]
    for i in ids:
        c = COL[i]
        r.append([ALC.get(m, m), c['n'], c['cp'], 'Sí' if cps.count(c['cp']) > 1 else 'No', len(ids), col_n[i], round(col_km[i], 3), i])
hojas.append(('Colonias homónimas', 'Colonias con el mismo nombre dentro de la misma alcaldía. En el buscador, los listados, los Excel y las fichas no se distinguen entre sí; cuando además repiten el código postal no hay forma de diferenciarlas.',
              ['Alcaldía', 'Colonia', 'Código postal', '¿Repite también el código postal?', 'Registros con ese nombre', 'Frentes', 'Km de frente', 'Identificador en el catálogo'], r))

r = [[ALC.get(c['m'], c['m']), c['n'], c['cp'], str(c['cp']).zfill(5), i] for i, c in enumerate(COL) if c.get('n') and len(str(c.get('cp', ''))) == 4]
r.sort()
hojas.append(('CP de cuatro dígitos', 'Códigos postales guardados con cuatro dígitos: se perdió el cero inicial (por ejemplo, 2900 en lugar de 02900). La columna «Código propuesto» solo antepone el cero.',
              ['Alcaldía', 'Colonia', 'Código postal en el catálogo', 'Código propuesto', 'Identificador en el catálogo'], r))

r = [[ALC.get(c['m'], c['m']), c['n'], c['cp'], c.get('pob', 0), col_n[i], i] for i, c in enumerate(COL) if c.get('n') and not c.get('ids')]
r.sort()
hojas.append(('Colonias sin índice social', 'Colonias sin Índice de Desarrollo Social asociado: su ficha no muestra el dato social de la unidad territorial.',
              ['Alcaldía', 'Colonia', 'Código postal', 'Población', 'Frentes', 'Identificador en el catálogo'], r))

ABREV = r'\b(ampl|fracc|frac|u h|uh|unid|hab|sn|sta|sto|gral|lic|ing|dr|prof|av|col|bo|pblo|secc|sec|ej|ctm|infonavit|fovissste|issste|rdcial|resid|cond|conj|mz|smz|lt|nte|sur|ote|pte|1a|2a|3a|1ra|2da|3ra|i|ii|iii|iv|v|vi)\b'
r = []
for i, c in enumerate(COL):
    n = c.get('n') or ''
    if not n:
        continue
    obs = []
    if '  ' in n or n != n.strip():
        obs.append('espacios dobles o al borde')
    ab = sorted(set(re.findall(ABREV, norm(n))) - {'sur', 'i', 'v'})
    if ab or '.' in n:
        obs.append('posible abreviatura o numeral: ' + ', '.join(ab) if ab else 'contiene punto')
    if re.search(r'["“”\'%#*]', n):
        obs.append('caracteres no usuales')
    if obs:
        r.append([ALC.get(c['m'], c['m']), n, c['cp'], '; '.join(obs), col_n[i], i])
r.sort()
hojas.append(('Colonias · formato', 'Nombres de colonia con posibles abreviaturas, numerales, espacios dobles o caracteres no usuales. La lista de abreviaturas es amplia a propósito: incluye numerales romanos y puntos cardinales abreviados para que el SIA decida cuáles se desarrollan.',
              ['Alcaldía', 'Colonia', 'Código postal', 'Observación', 'Frentes', 'Identificador en el catálogo'], r))

sin_acento = sum(1 for c in COL if c.get('n') and re.search(r'[áéíóúÁÉÍÓÚ]', c['n']))

# ---------- calles ----------
g = defaultdict(list)
for i, n in enumerate(NOM):
    if n and nom_n[i]:
        g[norm(n)].append(i)
r = []
k = 0
for clave, ids in sorted(g.items()):
    if len(ids) < 2 or not clave:
        continue
    k += 1
    mayor = max(ids, key=lambda i: nom_n[i])
    for i in sorted(ids, key=lambda i: -nom_n[i]):
        r.append([k, NOM[i], nom_n[i], round(nom_km[i], 3), 'más usada' if i == mayor else '', ', '.join(sorted(nom_alc[i]))])
hojas.append(('Calles · variantes', 'Nombres de calle que solo difieren en acentos, mayúsculas, espacios o puntuación. Cada grupo debería quedar con una sola escritura; se marca la más usada como referencia, no como propuesta definitiva.',
              ['Grupo', 'Nombre en el catálogo', 'Frentes', 'Km de frente', 'Referencia', 'Alcaldías donde aparece'], r))
grupos_calle = k

r = []
for i, n in enumerate(NOM):
    if not nom_n[i]:
        continue
    obs = []
    if re.search(r'["“”]', n):
        obs.append('comillas residuales')
    if re.match(r"^[^0-9A-Za-zÁÉÍÓÚÑáéíóúñ]", n):
        obs.append('inicia con un símbolo')
    if '  ' in n or n != n.strip():
        obs.append('espacios dobles o al borde')
    if norm(n) in ('prueba', 'test', 'ninguno', 'sin nombre', 'sn', 's n', 'na', 'n a', 'x', 'xx', 'xxx'):
        obs.append('posible residuo de captura o marcador de «sin nombre»')
    if re.fullmatch(r'[0-9 ]+', n):
        obs.append('solo números')
    if obs:
        r.append([n, '; '.join(obs), nom_n[i], round(nom_km[i], 3), ', '.join(sorted(nom_alc[i]))])
r.sort(key=lambda x: (x[1], x[0]))
hojas.append(('Calles · formato', 'Nombres de calle con comillas residuales, símbolos al inicio, solo números o que parecen un residuo de captura. Se listan para contrastar con la capa de frentes de INEGI; ninguno se corrigió.',
              ['Nombre en el catálogo', 'Observación', 'Frentes', 'Km de frente', 'Alcaldías donde aparece'], r))

r = []
for i, (mun, prio, name, col, ln, gc, x, y) in enumerate(FR):
    if norm(NOM[name]) == 'prueba':
        c = COL[col] if 0 <= col < len(COL) else {}
        r.append([i, NOM[name], ALC_I[mun], c.get('n', ''), META['prio'][prio], ln, round(y, 5), round(x, 5)])
hojas.append(('Calle «Prueba»', 'Frentes cuyo nombre de calle es «Prueba» (auditoría H-101). Puede ser un nombre real o un residuo de captura en la fuente: se contrasta con la capa de frentes de INEGI. Las coordenadas son las del primer vértice del frente.',
              ['Número de frente', 'Nombre', 'Alcaldía', 'Colonia', 'Prioridad', 'Longitud (m)', 'Latitud', 'Longitud'], r))

# ---------- vialidades primarias ----------
nomen = META['vp']['nomenclat']
vp_n = defaultdict(int)
vp_km = defaultdict(float)
for nom, nombre, mun, ln in VP:
    vp_n[nom] += 1
    vp_km[nom] += ln / 1000
r = []
gv = defaultdict(list)
for i, n in enumerate(nomen):
    if vp_n[i]:
        gv[re.sub(r'^(avenida|av|calzada|calz|boulevard|blvd|bulevar) ', '', norm(n))].append(i)
for i, n in enumerate(nomen):
    if vp_n[i] and not n.strip():
        r.append(['(vacío)', 'nombre vacío', vp_n[i], round(vp_km[i], 3)])
for clave, ids in sorted(gv.items()):
    if len(ids) > 1 and clave:
        for i in ids:
            r.append([nomen[i], 'misma vialidad con dos escrituras: ' + ' / '.join(nomen[j] for j in ids), vp_n[i], round(vp_km[i], 3)])
hojas.append(('Vialidades primarias', 'Nomenclaturas de vialidades primarias vacías o escritas de dos maneras (con y sin «Av.», «Avenida», «Calzada»). El agrupamiento es automático: revisar que cada par sea en efecto la misma vialidad.',
              ['Nomenclatura en el catálogo', 'Observación', 'Partes', 'Km sobre el eje'], r))

# ---------- Excel ----------
GUINDA, GRIS = '9D2148', '55585A'
wb = Workbook()
ws = wb.active
ws.title = 'Léeme'
lineas = [('Catálogos de colonias y de calles: listas para homologación', True),
          ('Herramienta «Modelo de priorización de reforestación urbana» · Secretaría del Medio Ambiente de la Ciudad de México · Sistema de Información Ambiental', False),
          ('', False),
          ('Qué es. Listas de registros de los catálogos que conviene revisar en la fuente de cada capa. Atiende los hallazgos H-029 y H-101 de la auditoría integral del 2 de octubre de 2026.', False),
          ('Qué no es. No es una corrección: la herramienta no modifica los catálogos. Cada lista se entrega para que el SIA decida y corrija en la capa de origen; después se regeneran los datos.', False),
          ('Origen. Datos publicados de la herramienta (%s frentes, %s colonias con nombre, %s nombres de calle en uso). Generado con 03_procesamiento_datos/reporte_catalogos.py.'
           % (format(N, ','), format(sum(1 for c in COL if c.get('n')), ','), format(sum(1 for i in range(len(NOM)) if nom_n[i]), ',')), False),
          (('Nota. Solo %s nombres de colonia llevan acento; el resto se escribe sin acentos, a diferencia del catálogo de calles.' % format(sin_acento, ',')) if sin_acento
           else 'Nota. Ningún nombre de colonia lleva acento (por ejemplo, «Ampliacion»), a diferencia del catálogo de calles, que sí los usa.', False),
          ('', False), ('Contenido', True)]
for t, b in lineas:
    ws.append([t])
    ws.cell(ws.max_row, 1).font = Font(bold=b, color=GUINDA if b else '000000', size=13 if b and ws.max_row == 1 else 11)
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
ws.append(['Hoja', 'Registros', 'Qué contiene'])
for c in range(1, 4):
    ws.cell(ws.max_row, c).font = Font(bold=True, color='FFFFFF')
    ws.cell(ws.max_row, c).fill = PatternFill('solid', fgColor=GUINDA)
for titulo, desc, enc, filas in hojas:
    ws.append([titulo, len(filas), desc])
    ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 110
for i in range(1, 9):
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
    ws.row_dimensions[i].height = 34 if i in (4, 5, 6, 7) else 18
for titulo, desc, enc, filas in hojas:
    h = wb.create_sheet(re.sub(r'[«»·]', '', titulo).replace('  ', ' ')[:31])
    h.append([desc])
    h.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(enc))
    h.cell(1, 1).alignment = Alignment(wrap_text=True, vertical='top')
    h.cell(1, 1).font = Font(italic=True, color=GRIS)
    h.row_dimensions[1].height = 48
    h.append(enc)
    for c in range(1, len(enc) + 1):
        h.cell(2, c).font = Font(bold=True, color='FFFFFF')
        h.cell(2, c).fill = PatternFill('solid', fgColor=GUINDA)
        h.cell(2, c).alignment = Alignment(wrap_text=True, vertical='center')
    for f in filas:
        h.append(f)
    for c in range(1, len(enc) + 1):
        ancho = max([len(str(enc[c - 1]))] + [len(str(f[c - 1])) for f in filas[:400]])
        h.column_dimensions[get_column_letter(c)].width = min(60, max(12, ancho + 2))
    h.freeze_panes = 'A3'
    h.auto_filter.ref = 'A2:%s%d' % (get_column_letter(len(enc)), max(2, len(filas) + 2))
os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
wb.save(SALIDA)
print('listas en', os.path.relpath(SALIDA, RAIZ))
for titulo, desc, enc, filas in hojas:
    print('  %-28s %6d' % (titulo, len(filas)))
print('  grupos de calles con variantes: %d' % grupos_calle)
